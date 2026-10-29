"""Programmatic document conversion with original embedded media and inline references."""

from __future__ import annotations
import base64
import io
import mimetypes
import re
from pathlib import Path
from urllib.parse import quote, urlparse, unquote
from urllib.request import urlopen
from collections import Counter
from zipfile import ZipFile
import xml.etree.ElementTree as ET


def require_text(source, markdown, label):
    """Reject missing source words/numbers, including repeated values."""
    from bs4 import BeautifulSoup
    import html
    visible = html.unescape(BeautifulSoup(markdown, "html.parser").get_text(" "))
    words = lambda text: Counter(re.findall(r"\w+", text, flags=re.UNICODE))
    missing = words(source) - words(visible)
    if missing:
        raise ValueError(f"Document conversion would lose text in {label}: {list(missing)[:8]}")


def markdown_html(html):
    from markdownify import MarkdownConverter

    class PreserveTables(MarkdownConverter):
        def convert_table(self, el, text, **kwargs):
            # Raw HTML tables are valid Markdown and retain spans and cell structure.
            from bs4 import BeautifulSoup
            table = BeautifulSoup(str(el), "html.parser")
            for node in list(table.find_all(["img", "a"])):
                link = node.get("src") if node.name == "img" else node.get("href")
                if link and link.startswith("memory://"):
                    label = node.get("alt") or node.get_text(" ") or "Media"
                    node.replace_with(f"[{label}]({link})")
            return "\n\n" + str(table) + "\n\n"

    return PreserveTables(heading_style="ATX", superscript="<sup>", subscript="<sub>").convert(html)



def pdf_table_markdown(table):
    from html import escape
    rows = table.rows
    rectangles = [cell for row in rows for cell in row.cells if cell is not None]
    xs = sorted({round(rect[i], 4) for rect in rectangles for i in (0, 2)})
    ys = sorted({round(rect[i], 4) for rect in rectangles for i in (1, 3)})
    values = table.extract()
    output = ["<table>"]
    for row_index, row in enumerate(rows):
        output.append("<tr>")
        for column, rect in enumerate(row.cells):
            if rect is None:
                continue
            x0, y0, x1, y1 = (round(v, 4) for v in rect)
            colspan = max(1, sum(x0 <= value < x1 for value in xs))
            rowspan = max(1, sum(y0 <= value < y1 for value in ys))
            value = values[row_index][column] or ""
            output.append(f'<td colspan="{colspan}" rowspan="{rowspan}">' + escape(value).replace("\n", "<br>") + "</td>")
        output.append("</tr>")
    output.append("</table>")
    return "\n".join(output)

def file_link(agent_id, session_id, item):
    return (
        f"memory://{quote(agent_id)}/{quote(session_id)}/{item['file_id']}"
        f"?filename={quote(item['filename'])}"
    )


class Documents:
    def __init__(self, store, agent_id, session_id):
        self.store, self.agent_id, self.session_id = store, agent_id, session_id
        self._parent = None

    def media(self, raw, filename, mime=None):
        mime = mime or mimetypes.guess_type(filename)[0] or "application/octet-stream"
        modality = mime.split("/")[0]
        kind = {"image": "image_url", "audio": "input_audio", "video": "video_url"}.get(
            modality
        )
        if kind is None:
            kind = "file"
        item = self.store.register_file(
            self.agent_id,
            raw,
            filename,
            kind,
            mime,
            siblings=[self._parent],
        )
        self._parent.setdefault("embedded_files", []).append(item)
        return f"[Media: {item['filename']}]({file_link(self.agent_id, self.session_id, item)})"

    def convert(self, item):
        self._parent = item
        path = self.store.instances / self.agent_id / "files" / item["filename"]
        suffix = path.suffix.lower()
        if suffix == ".docx":
            body = self.docx(path)
        elif suffix == ".pptx":
            body = self.pptx(path)
        elif suffix == ".pdf":
            body = self.pdf(path)
        elif suffix == ".xlsx":
            body = self.xlsx(path)
        elif suffix in {".html", ".htm", ".xhtml"}:
            body = self.html(path.read_bytes(), path.parent)
        elif suffix in {".txt", ".md", ".markdown", ".csv", ".tsv", ".json", ".xml", ".yaml", ".yml", ".log", ".rst"}:
            text = path.read_text(encoding="utf-8-sig")
            fence = "`" * max(3, 1 + max((len(m) for m in re.findall(r"`+", text)), default=0))
            body = text if suffix in {".md", ".markdown"} else f"{fence}{suffix[1:]}\n{text}\n{fence}"
        else:
            raise ValueError(f"No content-preserving document converter for {suffix or 'extensionless file'}")
        if not isinstance(body, str) or not body.strip():
            raise ValueError(
                f"Document contains no convertible content: {item['filename']}"
            )
        return (
            f"[Original document: {item['filename']}]({file_link(self.agent_id, self.session_id, item)})\n\n"
            + body
        )

    def docx(self, path):
        w = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
        metadata = []
        with ZipFile(path) as package:
            names = package.namelist()
            unsupported = [n for n in names if n.startswith(("word/charts/", "word/diagrams/", "word/embeddings/")) and not n.endswith("/")]
            if unsupported:
                raise ValueError(f"Cannot preserve Word embedded objects: {unsupported[:3]}")
            parts = [n for n in names if re.fullmatch(r"word/(document|header[0-9]+|footer[0-9]+|footnotes|endnotes|comments)\.xml", n)]
            text = []
            for name in parts:
                root = ET.fromstring(package.read(name))
                if name == "word/comments.xml":
                    for comment in root:
                        metadata.append("Comment " + comment.get(w + "id", "") + ": " + comment.get(w + "author", "") + " " + comment.get(w + "date", ""))
                for element in root.iter():
                    local = element.tag.rsplit("}", 1)[-1]
                    if local in {"oMath", "oMathPara", "del", "ins", "altChunk", "ruby"}:
                        raise ValueError(f"Cannot preserve Word element {local} in {name}")
                parents = {child: parent for parent in root.iter() for child in parent}
                paragraphs = {node: [] for node in root.iter(w + "p")}
                for node in root.iter(w + "t"):
                    owner = parents.get(node)
                    while owner is not None and owner.tag != w + "p":
                        owner = parents.get(owner)
                    if owner is not None:
                        paragraphs[owner].append(node.text or "")
                text.extend("".join(pieces) for pieces in paragraphs.values())
            # Headers and footers are omitted by Mammoth's ordinary body conversion.
            headers, footers = [], []
            for name in parts:
                if not re.fullmatch(r"word/(header|footer)[0-9]+\.xml", name):
                    continue
                root = ET.fromstring(package.read(name))
                if not list(root):
                    continue
                document = ET.Element(w + "document")
                body = ET.SubElement(document, w + "body")
                body.extend(list(root))
                rels = "word/_rels/" + Path(name).name + ".rels"
                empty_rels = b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"/>'
                stream = io.BytesIO()
                with ZipFile(stream, "w") as output:
                    for entry in names:
                        if entry not in {"word/document.xml", "word/_rels/document.xml.rels"}:
                            output.writestr(entry, package.read(entry))
                    output.writestr("word/document.xml", ET.tostring(document))
                    output.writestr("word/_rels/document.xml.rels", package.read(rels) if rels in names else empty_rels)
                stream.seek(0)
                converted = self._docx(stream)
                (headers if "header" in name else footers).append(f"## {Path(name).stem}\n\n" + converted)
        body = "\n\n".join([*headers, self._docx(path), *footers, *metadata])
        require_text(" ".join(text), body, path.name)
        return body

    def _docx(self, path):
        import mammoth
        from zipfile import ZipFile
        import hashlib
        import xml.etree.ElementTree as ET

        names = {}
        with ZipFile(path) as original:
            for name in original.namelist():
                if name.startswith("word/media/"):
                    names[hashlib.sha256(original.read(name)).hexdigest()] = Path(
                        name
                    ).name
            # Mammoth handles image positions. Insert audio/video links into the same
            # XML positions before conversion; do not modify the original document.
            xml = ET.fromstring(original.read("word/document.xml"))
            rels = ET.fromstring(original.read("word/_rels/document.xml.rels"))
            targets = {r.attrib["Id"]: r.attrib.get("Target", "") for r in rels}
            w = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
            parent_map = {child: parent for parent in xml.iter() for child in parent}
            audio_links = {}
            for element in list(xml.iter()):
                refs = [
                    value
                    for key, value in element.attrib.items()
                    if key.endswith(("}embed", "}link", "}id")) and value in targets
                ]
                for ref in refs:
                    target = targets[ref]
                    mime = mimetypes.guess_type(target)[0] or ""
                    if mime.split("/")[0] not in {"audio", "video"}:
                        continue
                    import posixpath

                    name = posixpath.normpath("word/" + target)
                    link = self.media(
                        original.read(name), Path(name).name, mime
                    )
                    owner = element
                    while owner is not None and owner.tag not in {w + "r", w + "p"}:
                        owner = parent_map.get(owner)
                    if owner is None:
                        raise ValueError(
                            "Cannot preserve the position of embedded document media"
                        )
                    run = (
                        owner if owner.tag == w + "r" else ET.SubElement(owner, w + "r")
                    )
                    token = "JARVISMEDIATOKEN" + str(len(audio_links))
                    audio_links[token] = link
                    ET.SubElement(run, w + "t").text = token
                    parent = parent_map.get(element)
                    if parent is not None and element in list(parent):
                        parent.remove(element)
            converted = io.BytesIO()
            with ZipFile(converted, "w") as output:
                for name in original.namelist():
                    data = original.read(name)
                    if name == "word/document.xml":
                        data = ET.tostring(xml)
                    elif name == "word/styles.xml":
                        styles = ET.fromstring(data)
                        for identifier, kind in (("CommentText", "paragraph"), ("CommentReference", "character")):
                            if not any(style.get(w + "styleId") == identifier for style in styles):
                                style = ET.SubElement(styles, w + "style", {w + "type": kind, w + "styleId": identifier})
                                ET.SubElement(style, w + "name", {w + "val": identifier})
                        data = ET.tostring(styles)
                    output.writestr(name, data)
        converted.seek(0)

        def image(image):
            with image.open() as stream:
                raw = stream.read()
            name = names.get(
                hashlib.sha256(raw).hexdigest(),
                "image" + (mimetypes.guess_extension(image.content_type) or ".bin"),
            )
            link = self.media(
                raw, name, image.content_type
            )
            return {"src": re.search(r"\]\((.*)\)$", link).group(1)}

        result = mammoth.convert_to_html(
            converted, convert_image=mammoth.images.img_element(image),
            style_map="\n".join(["comment-reference => sup", "p[style-name='header'] => p:fresh", "p[style-name='footer'] => p:fresh", "p[style-name='CommentText'] => p:fresh", "r[style-name='CommentReference'] => sup"]),
        )
        if result.messages:
            raise ValueError("Word conversion cannot guarantee preservation: " + "; ".join(m.message for m in result.messages))
        markdown = markdown_html(result.value)
        for token, link in audio_links.items():
            markdown = markdown.replace(token, link)
        return markdown

    def pptx(self, path):
        from markitdown.converters import PptxConverter
        from markitdown import StreamInfo
        from pptx.enum.shapes import MSO_SHAPE_TYPE

        documents = self

        class ArchivePptxConverter(PptxConverter):
            def _is_picture(self, shape):
                return (
                    super()._is_picture(shape)
                    or shape.shape_type == MSO_SHAPE_TYPE.MEDIA
                )

            def _convert_table_to_markdown(self, table, **kwargs):
                from html import escape
                output = ["<table>"]
                for row in table.rows:
                    output.append("<tr>")
                    for cell in row.cells:
                        if cell.is_spanned:
                            continue
                        output.append(f'<td colspan="{cell.span_width}" rowspan="{cell.span_height}">' + escape(cell.text).replace("\n", "<br>") + "</td>")
                    output.append("</tr>")
                output.append("</table>")
                return "\n\n" + "\n".join(output) + "\n\n"

            def _convert_picture_to_markdown(self, shape, **kwargs):
                # Keep MarkItDown's text/table/chart/group/note conversion. Only
                # substitute binary media with durable references at their positions.
                links, seen = [], set()
                for node in shape._element.iter():
                    for key, value in node.attrib.items():
                        if (
                            not key.endswith(("}embed", "}link"))
                            or value not in shape.part.rels
                        ):
                            continue
                        rel = shape.part.rels[value]
                        if rel.is_external:
                            continue
                        part = rel.target_part
                        mime = part.content_type
                        name = str(part.partname)
                        if (
                            mime.split("/")[0] in {"image", "audio", "video"}
                            and name not in seen
                        ):
                            links.append(
                                documents.media(
                                    part.blob, Path(name).name, mime
                                )
                            )
                            seen.add(name)
                if not links:
                    raise ValueError("Cannot extract original PowerPoint media")
                descriptions = [node.get("descr") or node.get("title") for node in shape._element.iter() if node.tag.endswith("}cNvPr")]
                return "\n" + "\n".join([*links, *filter(None, descriptions)]) + "\n"

        from pptx import Presentation
        with ZipFile(path) as package:
            if any(n.startswith(("ppt/diagrams/", "ppt/embeddings/", "ppt/comments/")) for n in package.namelist()):
                raise ValueError("Cannot preserve PowerPoint SmartArt, comments or embedded objects")
        presentation = Presentation(path)
        source = []
        def inspect_shape(shape):
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                for child in shape.shapes:
                    inspect_shape(child)
                return
            if shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
                raise ValueError("Cannot preserve PowerPoint vector shape geometry")
            if shape.has_text_frame:
                source.append(shape.text or "")
            if shape.has_table:
                source.extend(cell.text for row in shape.table.rows for cell in row.cells)
            if shape.has_chart:
                source.extend(node.text or "" for node in shape.chart._chartSpace.iter() if node.tag.rsplit("}", 1)[-1] in {"t", "v"})
            if not (shape.has_text_frame or shape.has_table or shape.has_chart or shape.shape_type in {MSO_SHAPE_TYPE.PICTURE, MSO_SHAPE_TYPE.MEDIA}):
                raise ValueError(f"Cannot preserve PowerPoint shape: {shape.shape_type}")
        for slide in presentation.slides:
            for shape in slide.shapes:
                inspect_shape(shape)
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
                source.append(slide.notes_slide.notes_text_frame.text)
        with path.open("rb") as stream:
            markdown = ArchivePptxConverter().convert(stream, StreamInfo(extension=".pptx")).markdown
        require_text(" ".join(source), markdown, path.name)
        return markdown

    def pdf(self, path):
        import pymupdf

        output = []
        with pymupdf.open(path) as document:
            if document.needs_pass:
                raise ValueError("Encrypted PDF requires a password")
            for number, page in enumerate(document, 1):
                output.append(f"## Page {number}\n\n")
                import inspect
                table_options = {"use_layout": False} if "use_layout" in inspect.signature(pymupdf.table.find_tables).parameters else {}
                tables = page.find_tables(**table_options).tables
                if any(not any(pymupdf.Rect(table.bbox).contains(drawing["rect"]) for table in tables) for drawing in page.get_drawings()):
                    raise ValueError(f"Cannot preserve PDF vector graphics on page {number}")
                if any(item.get("smask", 0) for item in (document.extract_image(info.get("xref")) for info in page.get_image_info(xrefs=True) if info.get("xref"))):
                    raise ValueError(f"Cannot preserve PDF image masks on page {number}")
                entries = [
                    (table.bbox[1], table.bbox[0], pdf_table_markdown(table) + "\n\n")
                    for table in tables
                ]
                for block in page.get_text("blocks"):
                    if block[6] == 0 and not any(
                        pymupdf.Rect(t.bbox).contains(pymupdf.Rect(block[:4]))
                        for t in tables
                    ):
                        entries.append((block[1], block[0], block[4].strip() + "\n\n"))
                for info in page.get_image_info(xrefs=True):
                    xref = info.get("xref")
                    if xref:
                        image = document.extract_image(xref)
                        link = self.media(
                            image["image"],
                            f"image_{xref}.{image['ext']}",
                        )
                    else:
                        # Inline PDF images have no standalone encoded stream; PyMuPDF
                        # exposes their original decoded image, rather than a page render.
                        block = next(
                            (
                                b
                                for b in page.get_text("dict")["blocks"]
                                if b["type"] == 1
                                and tuple(b["bbox"]) == tuple(info["bbox"])
                            ),
                            None,
                        )
                        if block is None:
                            raise ValueError("Cannot extract an inline PDF image")
                        link = self.media(
                            block["image"],
                            f"inline_{number}_{info['number']}.{block['ext']}",
                        )
                    entries.append((info["bbox"][1], info["bbox"][0], link + "\n\n"))
                for annotation in page.annots() or []:
                    if annotation.type[0] == pymupdf.PDF_ANNOT_FILE_ATTACHMENT:
                        info = annotation.file_info
                        name = (
                            info.get("ufilename")
                            or info.get("filename")
                            or "attachment.bin"
                        )
                        mime = mimetypes.guess_type(name)[0] or ""
                        link = self.media(annotation.get_file(), name, mime)
                        entries.append((annotation.rect.y0, annotation.rect.x0, link + "\n\n"))
                    elif annotation.info.get("content"):
                        entries.append((annotation.rect.y0, annotation.rect.x0, "\n\n" + annotation.info.get("title", "") + ": " + annotation.info["content"] + "\n\n"))
                    elif annotation.type[0] not in {pymupdf.PDF_ANNOT_LINK}:
                        raise ValueError(f"Cannot preserve PDF annotation on page {number}: {annotation.type}")
                for link in page.get_links():
                    destination = link.get("uri") or link.get("file") or str(link.get("page", ""))
                    entries.append((link["from"].y0, link["from"].x0, f"\n\n[Link]({destination})\n\n"))
                for field in page.widgets() or []:
                    entries.append((field.rect.y0, field.rect.x0, f"\n\n{field.field_name}: {field.field_value}\n\n"))
                converted = "".join(text for _, _, text in sorted(entries, key=lambda e: (e[0], e[1])))
                require_text(page.get_text(), converted, f"{path.name}, page {number}")
                output.append(converted)
        return "".join(output)

    def xlsx(self, path):
        import openpyxl
        from html import escape
        with ZipFile(path) as package:
            if any(n.startswith(("xl/charts/", "xl/embeddings/", "xl/pivotTables/")) for n in package.namelist()):
                raise ValueError("Cannot preserve spreadsheet charts, embedded objects or pivot tables")
        book = openpyxl.load_workbook(path, data_only=False)
        cached = openpyxl.load_workbook(path, data_only=True)
        output = []
        try:
            for sheet in book:
                output.append("# " + sheet.title + "\n\n")
                if sheet.sheet_state != "visible":
                    output.append(f"Sheet state: {sheet.sheet_state}\n\n")
                if sheet.conditional_formatting:
                    raise ValueError("Cannot preserve spreadsheet conditional formatting semantics")
                images = {}
                for image in sheet._images:
                    anchor = image.anchor._from
                    images.setdefault(anchor.row + 1, []).append(self.media(image._data(), Path(image.path).name))
                output.append("<table>\n")
                for row in sheet.iter_rows():
                    for link in images.pop(row[0].row, []):
                        output.append(f'<tr><td colspan="{sheet.max_column}">{link}</td></tr>\n')
                    cells = []
                    for cell in row:
                        value = "" if cell.value is None else str(cell.value)
                        text = escape(value).replace("\n", "<br>")
                        if cell.data_type == "f":
                            result = cached[sheet.title][cell.coordinate].value
                            if result is not None:
                                text += "<br>Cached value: " + escape(str(result))
                        if cell.hyperlink:
                            text += "<br>Link: " + escape(cell.hyperlink.target or cell.hyperlink.location or "")
                        if cell.comment:
                            text += "<br>Comment (" + escape(cell.comment.author or "") + "): " + escape(cell.comment.text).replace("\n", "<br>")
                        number_format = escape(cell.number_format, quote=True)
                        cells.append(f'<td data-cell="{cell.coordinate}" data-number-format="{number_format}">{text}</td>')
                    output.append("<tr>" + "".join(cells) + "</tr>\n")
                output.append("</table>\n\n")
                for links in images.values():
                    output.extend(link + "\n\n" for link in links)
                if sheet.merged_cells.ranges:
                    output.append("Merged cells: " + ", ".join(str(r) for r in sheet.merged_cells.ranges) + "\n\n")
        finally:
            book.close()
            cached.close()
        return "".join(output)

    def html(self, html, base):
        from bs4 import BeautifulSoup

        soup = BeautifulSoup(html, "html.parser")
        if soup.find(["svg", "canvas", "iframe", "object", "embed", "script", "track"]):
            raise ValueError("Cannot preserve HTML dynamic content, inline vector graphics or embedded objects")
        original_text = soup.get_text(" ")
        for i, node in enumerate(soup.find_all(["img", "audio", "video"])):
            source = node.get("src") or (node.find("source") or {}).get("src")
            if not source:
                raise ValueError("Embedded media has no source")
            if source.startswith("data:"):
                header, data = source.split(",", 1)
                mime = header[5:].split(";")[0]
                raw = base64.b64decode(data, validate=True)
                name = f"media_{i}" + (mimetypes.guess_extension(mime) or ".bin")
            elif source.startswith(("https://", "http://")):
                with urlopen(source, timeout=60) as response:
                    raw, mime = response.read(), response.headers.get_content_type()
                name = Path(urlparse(source).path).name
            else:
                path = base / unquote(source)
                raw = path.read_bytes()
                name, mime = path.name, mimetypes.guess_type(path.name)[0]
            link = self.media(raw, name, mime)
            anchor = soup.new_tag("a", href=re.search(r"\]\((.*)\)$", link).group(1))
            anchor.string = " ".join(filter(None, ["Media: " + name, node.get("alt"), node.get("title"), node.get_text(" ")]))
            node.replace_with(anchor)
        markdown = markdown_html(str(soup))
        require_text(original_text, markdown, "HTML document")
        return markdown
