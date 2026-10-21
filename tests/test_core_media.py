"""Core adaptation preserves originals, caches interpretations and honors actual routes."""
import base64
import copy
import io
import json
import unittest
from types import SimpleNamespace
from unittest.mock import Mock, patch

import test_runtime as runtime
from test_session_memory import PNG, FakeMemoryAPI
from jarvis.core.protocol import Event
from jarvis.infrastructure.context import MemoryStore
from jarvis.infrastructure.model_capabilities import ModelCapabilities
from jarvis.memory.serialization import summary_context
from test_subscription import _client, _ResponseStream


class CoreMediaTests(unittest.TestCase):
    def setUp(self):
        self.fixture = runtime.RuntimeTests()
        self.fixture.setUp()
        self.addCleanup(self.fixture.doCleanups)
        self.addCleanup(self.fixture.tearDown)
        self.store = MemoryStore(self.fixture.root / 'memory')
        self.manager = self.fixture.manager(runtime._Client([]), memory=self.store)
        self.manager.restoring_shutdown = True
        self.agent = self.manager.spawn_root(preset='main')
        self.manager.config = SimpleNamespace(subscription_enabled=False, base_url=None,
            memory_model_text='api/text', memory_model_text_image='api/image',
            memory_model_text_image_video='api/video', memory_model_text_image_audio_video='api/full')
        self.api = FakeMemoryAPI()
        self.manager.services['memory_client'] = self.api
        self.manager.model_capabilities = ModelCapabilities('test', ('text',))

    def attach(self, part):
        self.agent.history.append({'role': 'user', 'content': [{'type': 'text', 'text': 'Original event'}, part]})
        self.agent._persist_context()

    def image(self):
        return {'type': 'image_url', 'image_url': {'name': 'picture.png',
            'url': 'data:image/png;base64,' + base64.b64encode(PNG).decode()}}

    def test_unsupported_image_becomes_identified_text_and_is_cached(self):
        self.attach(self.image())
        request = self.agent._model_messages()
        part = request[1]['content'][1]
        self.assertEqual(part['type'], 'text')
        self.assertIn('file_000001', part['text'])
        self.assertIn('session_000001', part['text'])
        self.assertIn('main', part['text'])
        self.assertIn('Работа сохранена.', part['text'])
        self.agent._model_messages()
        self.assertEqual(len(self.api.calls), 1)
        context = self.store.load('main', 'main')['messages']
        self.assertEqual(context[0]['content'][1]['type'], 'image_url')
        self.assertEqual((self.store.instances / 'main/files/picture.png').read_bytes(), PNG)

    def test_supported_media_remains_native_without_interpretation(self):
        self.attach(self.image())
        request = self.agent._model_messages(capabilities=ModelCapabilities('vision', ('text', 'image')))
        self.assertEqual(request[1]['content'][1]['type'], 'image_url')
        self.assertEqual(self.api.calls, [])

    def test_original_unsupported_event_is_saved_at_enqueue(self):
        from jarvis.capabilities.api import input_part
        event = Event('tick.event', {}, parts=(input_part('image', 'image/png', base64.b64encode(PNG).decode(), name='raw.png'),))
        self.agent.enqueue(event)
        self.assertEqual(self.store.files('main')[0]['type'], 'image_url')
        self.assertEqual(self.api.calls, [])

    def test_new_event_arriving_during_interpretation_is_preserved(self):
        self.attach(self.image())
        original = self.api.create
        def complete(**options):
            self.agent.enqueue(Event('tick.event', {'text': 'arrived during API call'}))
            return original(**options)
        self.api.create = complete
        request = self.agent._model_messages()
        self.assertEqual(len(request), 2)
        live = self.store.load('main', 'main')['messages']
        self.assertEqual(len(live), 2)
        self.assertIn('arrived during API call', live[-1]['content'])
        self.assertTrue(live[0]['content'][1]['text_representation'])

    def test_document_interleaves_supported_native_media_and_interprets_only_when_needed(self):
        from docx import Document
        doc = Document()
        doc.add_paragraph('Before picture')
        doc.add_picture(io.BytesIO(PNG))
        doc.add_paragraph('After picture')
        stream = io.BytesIO()
        doc.save(stream)
        self.attach({'type': 'file', 'file': {'filename': 'report.docx',
            'file_data': 'data:application/octet-stream;base64,' + base64.b64encode(stream.getvalue()).decode()}})
        request = self.agent._model_messages(capabilities=ModelCapabilities('vision', ('text', 'image')))
        parts = request[1]['content']
        types = [p['type'] for p in parts]
        self.assertIn('image_url', types)
        self.assertNotIn('file', types)
        self.assertEqual(self.api.calls, [])
        self.assertIn('Before picture', parts[1]['text'])
        self.assertIn('After picture', parts[-1]['text'])
        request = self.agent._model_messages()
        self.assertTrue(all(p['type'] == 'text' for p in request[1]['content']))
        self.assertEqual(len(self.api.calls), 1)
        self.assertEqual(len(self.store.files('main')), 2)

    def test_subscription_fallback_rebuilds_from_originals_for_api_modalities(self):
        self.attach(self.image())
        auth = Mock()
        auth.credentials.return_value = {'access_token': 'test', 'account_id': 'test'}
        session = Mock()
        session.post.return_value = _ResponseStream(status=429, text='usage limit')
        router, api = _client(auth, session, api_capabilities=ModelCapabilities('api-model', ('text',)))
        self.addCleanup(router.close)
        api.chat.completions.create.return_value = iter([])
        history = copy.deepcopy(self.agent.history)
        prepare = lambda model, caps: self.agent._model_messages('test', capabilities=caps, history=history)
        stream = router.create(model='gpt-6-luna', messages=history, stream=True, _jarvis_messages=prepare)
        list(stream)
        stream.close()
        sent = api.chat.completions.create.call_args.kwargs
        self.assertNotIn('_jarvis_messages', sent)
        self.assertEqual(sent['messages'][1]['content'][1]['type'], 'text')
        self.assertIn('file_000001', sent['messages'][1]['content'][1]['text'])
        self.assertEqual(len(self.api.calls), 1)

    def test_audio_and_video_fallbacks_use_their_profiles_without_losing_originals(self):
        self.attach({'type': 'input_audio', 'input_audio': {'format': 'wav', 'data': base64.b64encode(b'audio').decode()}})
        self.attach({'type': 'video_url', 'video_url': {'name': 'clip.mp4', 'url': 'data:video/mp4;base64,' + base64.b64encode(b'video').decode()}})
        with patch('jarvis.core.media.subprocess.run', return_value=SimpleNamespace(stdout='{"streams": []}')):
            request = self.agent._model_messages()
        self.assertTrue(all(m['content'][1]['type'] == 'text' for m in request[1:]))
        self.assertEqual([call['model'] for call in self.api.calls], ['api/full', 'api/video'])
        self.assertEqual([f['type'] for f in self.store.files('main')], ['input_audio', 'video_url'])

    def test_failed_interpretation_preserves_original_and_can_be_retried(self):
        self.attach(self.image())
        self.api.fail_on = 1
        with self.assertRaises(RuntimeError):
            self.agent._model_messages()
        self.assertEqual(self.store.files('main')[0]['text_representation'], '')
        self.api.fail_on = None
        request = self.agent._model_messages()
        self.assertEqual(request[1]['content'][1]['type'], 'text')
        self.assertEqual(len(self.store.files('main')), 1)

    def test_summary_request_omits_shutdown_but_archive_retains_it(self):
        from jarvis.memory.archiver import Archiver
        self.agent.history.extend([
            {'role': 'user', 'content': 'Please finish'},
            {'role': 'assistant', 'content': json.dumps({'actions': [{'action_id': 'shutdown', 'call_id': 's', 'data': {}}]})},
        ])
        self.agent._persist_context()
        self.store.update_state(shutdown={'session_id': 'session_000001'})
        archiver = Archiver(self.store, self.manager.config, self.api, 'environment')
        archiver.summarize_instances()
        self.assertNotIn('shutdown', str(self.api.calls[0]['messages'][1:]))
        archived = self.store.root / '.archive-session_000001/instances/main/context.json'
        self.assertIn('shutdown', archived.read_text())

    def test_shutdown_formalities_are_removed_only_from_summary_input(self):
        history = [
            {'role': 'user', 'content': 'Please finish'},
            {'role': 'assistant', 'thinking': 'technical shutdown reasoning', 'content': json.dumps({'actions': [{'action_id': 'shutdown', 'call_id': 's', 'data': {}}]})},
            {'role': 'user', 'content': json.dumps({'event_id': 'call_result', 'call_id': 's', 'data': {'status': 'completed'}})},
        ]
        filtered = summary_context(history)
        self.assertEqual(filtered, history[:1])
        self.assertEqual(len(history), 3)
