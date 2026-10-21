"""Behavioral tests of durable sessions, native archives and restart checkpoints."""

import base64
import io
import json
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock, patch

import fixtures
import test_runtime as runtime
from jarvis.core.agent import Agent
from jarvis.core.protocol import CallResult, Event
from jarvis.infrastructure.config import Config
from jarvis.infrastructure.context import MemoryStore, hydrate_part, read_json
from jarvis.memory.archiver import Archiver
from jarvis.memory.documents import Documents
from jarvis.memory.reader import ArchiveReader
from jarvis.memory.serialization import chunks, serialize_context

PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+aXioAAAAASUVORK5CYII="
)


class CharTokenizer:
    def encode(self, text, **kwargs):
        return list(map(ord, text))

    def decode(self, values, **kwargs):
        return "".join(map(chr, values))


class FakeMemoryAPI:
    def __init__(self, fail_on=None):
        self.calls = []
        self.fail_on = fail_on
        self.chat = SimpleNamespace(completions=self)

    def create(self, **kwargs):
        self.calls.append(kwargs)
        if len(self.calls) == self.fail_on:
            raise RuntimeError("Simulated API outage")
        return SimpleNamespace(
            choices=[
                SimpleNamespace(
                    message=SimpleNamespace(content="# Account\nРабота сохранена.")
                )
            ]
        )


class SessionMemoryTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name)
        self.store = MemoryStore(self.root / "memory")

    def record(self, agent="main", messages=None):
        return {
            "agent_id": agent,
            "preset_id": "main" if agent == "main" else "worker",
            "parent_id": None if agent == "main" else "main",
            "messages": messages or [],
            "description": "A test instance",
        }

    def attachment(self, name="report.png", raw=PNG):
        return {
            "type": "image_url",
            "image_url": {
                "url": "data:image/png;base64," + base64.b64encode(raw).decode(),
                "name": name,
            },
        }

    def test_names_and_ids_are_local_durable_and_original_bytes_survive(self):
        messages = [
            {"role": "user", "content": [self.attachment()]},
            {"role": "user", "content": [self.attachment(raw=PNG + b"2")]},
        ]
        self.store.save(self.record(messages=messages))
        self.store.save(self.record(messages=messages))
        self.store.save(self.record("agent-00001", messages=messages[:1]))
        parts = read_json(self.store.instances / "main/context.json")
        self.assertEqual(
            [m["content"][0]["filename"] for m in parts], ["report.png", "report_1.png"]
        )
        self.assertEqual(parts[0]["content"][0]["file_id"], "file_000001")
        self.assertEqual(self.store.files("agent-00001")[0]["file_id"], "file_000001")
        self.assertEqual(
            (self.store.instances / "main/files/report.png").read_bytes(), PNG
        )
        self.assertEqual(len(self.store.files("main")), 2)
        restored = MemoryStore(self.store.root).load("main", "main")["messages"][0][
            "content"
        ][0]
        self.assertEqual(
            hydrate_part(restored, self.store.instances / "main")["image_url"]["url"],
            messages[0]["content"][0]["image_url"]["url"],
        )

    def test_context_is_the_only_file_metadata_and_manual_edits_survive_restart(self):
        self.store.save(self.record(messages=[{"role": "user", "content": [self.attachment()]}]))
        directory = self.store.instances / "main"
        self.assertFalse((directory / "files.json").exists())
        history = read_json(directory / "context.json")
        part = history[0]["content"][0]
        self.assertEqual(set(part), {"type", "file_id", "filename", "text_representation"})
        (directory / "files/report.png").rename(directory / "files/renamed.png")
        part.update(filename="renamed.png", text_representation="Manually corrected")
        self.store._write_json(directory / "context.json", history)
        restored = MemoryStore(self.store.root)
        record = restored.load("main", "main")
        restored.save(record)
        self.assertEqual(restored.files("main")[0], part)
        self.assertIn("data:image/png", hydrate_part(part, directory)["image_url"]["url"])

    def test_original_document_named_files_json_is_not_treated_as_a_registry(self):
        native = {"type": "file", "file": {"filename": "files.json",
                  "file_data": "data:application/json;base64," + base64.b64encode(b'{"important": true}').decode()}}
        self.store.save(self.record(messages=[{"role": "user", "content": [native]}]))
        restarted = MemoryStore(self.store.root)
        self.assertEqual((restarted.instances / "main/files/files.json").read_bytes(), b'{"important": true}')

    def test_legacy_index_migrates_embedded_media_into_context_and_is_removed(self):
        self.store.save(self.record())
        directory = self.store.instances / "main"
        child = {"type": "image_url", "file_id": "file_000002", "filename": "image.png",
                 "mime_type": "image/png", "source": "old", "sha256": "old", "native_options": {},
                 "text_representation": "Image"}
        document = {"type": "file", "file_id": "file_000001", "filename": "report.docx",
                    "text_representation": "Document", "agent_id": "main", "session_id": "session_000001",
                    "document_markdown": "[Media](memory://main/session_000001/file_000002?filename=image.png)"}
        self.store._write_json(directory / "context.json", [{"role": "user", "content": [document]}])
        self.store._write_json(directory / "files.json", {"files": [document, child]})
        restored = MemoryStore(self.store.root)
        self.assertFalse((directory / "files.json").exists())
        items = restored.files("main")
        self.assertEqual([f["file_id"] for f in items], ["file_000001", "file_000002"])
        self.assertNotIn("agent_id", items[0])
        self.assertEqual(set(items[1]), {"type", "file_id", "filename", "text_representation"})

    def test_document_archival_resumes_after_media_api_failure_without_new_ids(self):
        from docx import Document
        document = Document()
        document.add_paragraph("Before")
        document.add_picture(io.BytesIO(PNG))
        document.add_paragraph("After")
        stream = io.BytesIO()
        document.save(stream)
        native = {"type": "file", "file": {"filename": "report.docx",
                  "file_data": "data:application/octet-stream;base64," + base64.b64encode(stream.getvalue()).decode()}}
        self.store.save(self.record(messages=[{"role": "user", "content": [native]}]))
        self.store.update_state(shutdown={"session_id": "session_000001"})
        config = SimpleNamespace(memory_model_text="api/text", memory_model_text_image="api/image")
        with self.assertRaises(RuntimeError):
            Archiver(self.store, config, FakeMemoryAPI(fail_on=1), "env").interpret()
        before = self.store.files("main")
        self.assertEqual(len(before), 2)
        restarted = MemoryStore(self.store.root)
        archiver = Archiver(restarted, config, FakeMemoryAPI(), "env")
        archiver.interpret()
        self.assertEqual([f["file_id"] for f in restarted.files("main")], [f["file_id"] for f in before])
        archiver.summarize_instances()
        archiver.summarize_session()
        archiver.commit()
        directory = restarted.sessions / "session_000001/instances/main"
        self.assertFalse((directory / "files.json").exists())
        metadata = read_json(directory / "instance.json")
        self.assertNotIn("pending_turn", metadata)
        self.assertNotIn("capabilities_add", metadata)
        name, media = ArchiveReader(restarted.root, Mock()).get_file("main", "session_000001", before[1]["file_id"])
        self.assertEqual(name, before[1]["filename"])
        self.assertEqual(base64.b64decode(media["image_url"]["url"].split(",", 1)[1]), PNG)

    def test_prepared_context_transaction_is_recovered_after_a_write_failure(self):
        self.store.save(self.record(messages=[{"role": "user", "content": "before"}]))
        real = self.store._write_json

        def fail(path, value):
            if path.name == "instance.json":
                raise OSError("disk full")
            return real(path, value)

        with patch.object(self.store, "_write_json", side_effect=fail):
            with self.assertRaises(OSError):
                self.store.save(
                    self.record(messages=[{"role": "user", "content": "latest"}])
                )
        restored = MemoryStore(self.store.root).load("main", "main")
        self.assertEqual(restored["messages"][0]["content"], "latest")

    def test_compact_full_history_and_search_history_have_distinct_reasoning(self):
        history = [
            {
                "role": "user",
                "content": '{ "event_id": "speech_detected", "data": {"text": "Сделай игру"} }',
            },
            {
                "role": "assistant",
                "thinking": "Поручу разработку.",
                "content": '{"actions":[]}',
            },
            {
                "role": "user",
                "content": [
                    {"type": "text", "text": '{"event_id":"call_result","data":{}}'},
                    {
                        "type": "image_url",
                        "agent_id": "main",
                        "session_id": "session_000001",
                        "file_id": "file_000001",
                        "filename": "screen.png",
                        "text_representation": "Скриншот игры.",
                    },
                ],
            },
        ]
        full = serialize_context(history)
        self.assertIn(
            '<agent_start><thinking_start>Поручу разработку.<thinking_end>{"actions":[]}<agent_end>',
            full,
        )
        self.assertIn("<image_interpretation_start>", full)
        self.assertNotIn("\n", full)
        self.assertNotIn("Поручу", serialize_context(history, include_thinking=False))

    def test_token_chunks_cover_long_messages_overlap_and_keep_inclusive_indexes(self):
        history = [
            {"role": "user", "content": str(i) + " " + "a" * 400} for i in range(30)
        ]
        tokenizer = CharTokenizer()
        split = chunks(history, tokenizer)
        self.assertTrue(20 <= len(split) <= 40)
        self.assertEqual(split[0]["start"], 0)
        self.assertEqual(split[-1]["end"], 29)
        self.assertTrue(all(a["end"] >= b["start"] for a, b in zip(split, split[1:])))
        long = chunks(
            [{"role": "user", "content": "x" * 10000}], tokenizer, max_tokens=300
        )
        self.assertGreater(len(long), 1)
        self.assertTrue(all(c["start"] == c["end"] == 0 for c in long))
        self.assertLess(
            len(chunks([{"role": "user", "content": "short"}], tokenizer)), 20
        )

    def test_reader_uses_original_message_slices_and_only_completed_sessions(self):
        directory = self.store.sessions / "session_000001/instances/main"
        directory.mkdir(parents=True)
        history = [{"role": "user", "content": str(i)} for i in range(400)]
        self.store._write_json(directory / "context.json", history)
        reader = ArchiveReader(self.store.root, Mock())
        self.assertEqual(reader.list_sessions("main"), {"sessions": []})
        (directory.parent.parent / "summary.md").write_text("Session")
        (directory / "summary.md").write_text("Instance")
        self.store._write_json(directory / "instance.json", self.record())
        result = reader.get_context("main", "session_000001", 120, 243)
        self.assertEqual(result["total"], 37)
        self.assertIn("<event_start>120<event_end>", result["context"])
        self.assertNotIn("<event_start>157<event_end>", result["context"])
        self.assertEqual(
            reader.get_context_info("main", "session_000001"),
            {"messages": 400, "characters": len(serialize_context(history))},
        )
        with self.assertRaises(ValueError):
            reader.get_context("../main", "session_000001")

    def test_summary_search_streams_all_sessions_and_returns_historical_preset(self):
        for n in range(1, 5):
            directory = self.store.sessions / f"session_{n:06d}/instances/main"
            directory.mkdir(parents=True)
            (directory.parent.parent / "summary.md").write_text("Session")
            (directory / "summary.md").write_text(str(n))
            self.store._write_json(directory / "context.json", [])
            self.store._write_json(
                directory / "instance.json", {"preset_id": "historical"}
            )
        reranker = Mock()
        reranker.score.side_effect = lambda q, candidates: [int(candidates[0]) / 10]
        result = ArchiveReader(self.store.root, reranker).search("main", "topic", 2)
        self.assertEqual(
            [r["session_id"] for r in result["matches"]],
            ["session_000004", "session_000003"],
        )
        self.assertEqual(reranker.score.call_count, 4)
        self.assertTrue(all(r["preset_id"] == "historical" for r in result["matches"]))

    def test_model_selection_requires_explicit_profiles_and_detects_video_audio(self):
        config = SimpleNamespace(
            **{
                "memory_model_" + name: "api/" + name
                for name in [
                    "text",
                    "text_image",
                    "text_image_video",
                    "text_image_audio_video",
                ]
            }
        )
        archiver = Archiver(self.store, config, Mock(), "environment")
        self.assertEqual(archiver.model({"text", "image"}), "api/text_image")
        self.assertEqual(archiver.model({"audio"}), "api/text_image_audio_video")
        part = {"type": "video_url", "filename": "clip.mp4"}
        with patch(
            "jarvis.memory.archiver.subprocess.run",
            return_value=SimpleNamespace(stdout='{"streams":[{"codec_type":"audio"}]}'),
        ):
            self.assertIn("audio", archiver.modalities([part], self.root))
        with self.assertRaises(ValueError):
            Archiver(self.store, SimpleNamespace(), Mock(), "env").model({"text"})

    def test_document_summary_interleaves_original_media_without_caption_injection(
        self,
    ):
        self.store.save(self.record())
        image = self.store.register_file(
            "main", PNG, "embedded.png", "image_url", "image/png"
        )
        self.store.save(self.record(messages=[{"role": "user", "content": [image]}]))
        markdown = f"Before [Media](memory://main/session_000001/{image['file_id']}?filename=embedded.png) After"
        archiver = Archiver(self.store, SimpleNamespace(), Mock(), "env")
        parts = archiver.document_parts(
            {
                "type": "file",
                "file_id": "file_999999",
                "document_markdown": markdown,
                "text_representation": "ignored",
            },
            self.store.instances / "main",
        )
        self.assertEqual([p["type"] for p in parts], ["text", "image_url", "text"])
        self.assertTrue(parts[0]["text"].startswith("Before"))
        self.assertEqual(parts[2]["text"], " After")


class ShutdownTests(unittest.TestCase):
    def setUp(self):
        self.fixture = runtime.RuntimeTests()
        self.fixture.setUp()
        self.addCleanup(self.fixture.doCleanups)
        self.addCleanup(self.fixture.tearDown)
        self.store = MemoryStore(self.fixture.root / "memory")
        self.api = FakeMemoryAPI()
        self.manager = self.make_manager()
        with patch.object(Agent, "start", lambda agent: agent._contract()):
            self.main = self.manager.spawn_root(preset="main")

    def make_manager(self):
        farewell = '{"actions":[{"action_id":"reply","call_id":"farewell","data":{"text":"Сессия завершена."}}]}'
        manager = self.fixture.manager(runtime._Client([farewell]), memory=self.store)
        manager.actions.unregister_owner("core:speech")
        manager.config = Config(
            "test",
            None,
            "test",
            self.fixture.root,
            self.fixture.root,
            memory_model_text="api/text",
            memory_model_text_image="api/image",
            memory_model_text_image_video="api/video",
            memory_model_text_image_audio_video="api/all",
        )
        manager.services["memory_api_client"] = self.api
        return manager

    def stop_session(self):
        self.manager.session_shutdown.start("shutdown-1")
        self.manager.session_shutdown.thread.join(timeout=10)
        self.assertFalse(self.manager.session_shutdown.thread.is_alive())

    def test_shutdown_archives_interrupted_actions_and_farewell_is_outside_history(
        self,
    ):
        self.main.enqueue(Event("tick.event", {"text": "Do work"}, target="main"))
        self.main.history.append(
            {
                "role": "assistant",
                "thinking": "Private thought",
                "content": '{"actions":[{"action_id":"say","call_id":"running","data":{"text":"hello"}}]}',
            }
        )
        self.main._persist_context()
        self.manager.results.begin(
            agent_id="main",
            call_id="running",
            action_id="say",
            result_schema={"properties": {}},
        )
        self.stop_session()
        archive = self.store.sessions / "session_000001/instances/main"
        context = read_json(archive / "context.json")
        self.assertTrue(
            any("Action interrupted by session shutdown." in str(m) for m in context)
        )
        self.assertFalse(
            any(
                "farewell" in str(m) or '"status":"completed"' in str(m)
                for m in context
            )
        )
        self.assertEqual(read_json(self.store.instances / "main/context.json"), [])
        self.assertEqual(self.main.history[1:], [])
        self.assertTrue(self.manager.completed.is_set())
        self.assertFalse(self.main.enqueue(Event("tick.event", {"text": "late"})))
        self.assertEqual(len(self.api.calls), 2)
        self.assertTrue(
            any(m["role"] == "assistant" for m in self.api.calls[0]["messages"])
        )
        self.assertIn("Private thought", str(self.api.calls[0]["messages"]))
        self.assertEqual(self.store.state()["session_id"], "session_000002")

    def test_api_failure_keeps_live_context_and_restart_resumes_without_redoing_summary(
        self,
    ):
        self.api.fail_on = 2  # individual summary succeeds, combined request fails
        self.main.enqueue(Event("tick.event", {"text": "Preserve me"}))
        self.stop_session()
        self.assertFalse(self.manager.completed.is_set())
        self.assertEqual(self.store.state()["shutdown"]["stage"], "summary")
        self.assertFalse((self.store.sessions / "session_000001").exists())
        self.assertTrue(read_json(self.store.instances / "main/context.json"))
        self.manager.shutdown()
        self.api = FakeMemoryAPI()
        self.manager = self.make_manager()
        self.manager.restoring_shutdown = True
        self.manager.frozen.set()
        self.main = self.manager.restore(preset="main")
        self.stop_session()
        self.assertEqual(len(self.api.calls), 1)
        self.assertTrue((self.store.sessions / "session_000001/summary.md").exists())
        self.assertTrue(self.manager.completed.is_set())

    def test_deleted_instance_is_archived_but_not_restored_in_next_session(self):
        with patch.object(Agent, "start", lambda agent: agent._contract()):
            child_id = self.manager.spawn(parent_id="main", preset_id="worker")[
                "agent_id"
            ]
        child = self.manager.require_agent(child_id)
        child.enqueue(Event("tick.event", {"text": "Child activity"}))
        child.semantic_memory.write("Own fact")
        self.manager.delete(agent_id=child_id)
        self.assertNotIn(child_id, self.manager.agents)
        self.assertTrue(self.store.load("worker", child_id)["deleted"])
        self.stop_session()
        archived = self.store.sessions / "session_000001/instances" / child_id
        self.assertTrue((archived / "context.json").exists())
        self.assertIn("instance_deleted", (archived / "context.json").read_text())
        self.assertFalse((self.store.instances / child_id).exists())
        self.manager.shutdown()
        self.manager = self.make_manager()
        self.assertFalse(self.manager.session_shutdown.pending)
        with patch.object(Agent, "start", lambda agent: agent._contract()):
            main = self.manager.restore(preset="main")
            new_id = self.manager.spawn(parent_id="main", preset_id="worker")[
                "agent_id"
            ]
        self.assertEqual(main.agent_id, "main")
        self.assertNotEqual(new_id, child_id)

    def test_semantic_memories_are_individual_and_survive_shutdown(self):
        with patch.object(Agent, "start", lambda agent: agent._contract()):
            child_id = self.manager.spawn(parent_id="main", preset_id="worker")[
                "agent_id"
            ]
        child = self.manager.require_agent(child_id)
        self.main.semantic_memory.write("Main fact")
        child.semantic_memory.write("Child fact")
        child._contract()
        self.assertIn("Child fact", child.history[0]["content"])
        self.assertNotIn("Main fact", child.history[0]["content"])
        self.assertIn("semantic_memory.write", child._contract()[0])
        self.stop_session()
        self.assertEqual(
            child.semantic_memory.snapshot()["entries"][0]["content"], "Child fact"
        )
        self.assertEqual(
            self.main.semantic_memory.snapshot()["entries"][0]["content"], "Main fact"
        )

    def test_interpretation_failure_is_resumable_and_keeps_original(self):
        image = {
            "type": "image_url",
            "image_url": {
                "url": "data:image/png;base64," + base64.b64encode(PNG).decode(),
                "name": "original.png",
            },
        }
        self.main.enqueue(CallResult("image", {}, "main", parts=()))
        self.main.history.append({"role": "user", "content": [image]})
        self.main._persist_context()
        self.api.fail_on = 1
        self.stop_session()
        self.assertEqual(self.store.state()["shutdown"]["stage"], "interpret")
        self.assertEqual(
            (self.store.instances / "main/files/original.png").read_bytes(), PNG
        )
        self.assertFalse((self.store.sessions / "session_000001").exists())

    def test_shutdown_is_main_only_and_memory_is_indivisible(self):
        with patch.object(Agent, "start", lambda agent: agent._contract()):
            child_id = self.manager.spawn(parent_id="main", preset_id="worker")[
                "agent_id"
            ]
        child = self.manager.require_agent(child_id)
        self.assertNotIn("shutdown", child._contract()[0])
        with self.assertRaises(ValueError):
            self.manager.capabilities.partition(["memory.search"])
        self.manager.enable_module(child_id, "memory")
        self.assertEqual(
            len([n for n in child._contract()[0] if n.startswith("memory.")]), 7
        )

    def test_shutdown_forcibly_stops_a_real_running_capability_process(self):
        from jarvis.core.protocol import ActionRequest

        marker = self.fixture.root / "started.txt"
        self.fixture.write_action(
            "slow",
            runtime.EMPTY_SCHEMA,
            {
                "type": "object",
                "properties": {"status": {"type": "string"}},
                "required": ["status"],
            },
            "    import time\n    from pathlib import Path\n    Path("
            + repr(str(marker))
            + ").write_text('running')\n    time.sleep(60)\n    return {'status':'done'}\n",
        )
        self.manager.enable_action("main", "slow")
        action = {"action_id": "slow", "call_id": "slow-call", "data": {}}
        self.main.history.append(
            {"role": "assistant", "content": json.dumps({"actions": [action]})}
        )
        self.main._persist_context()
        self.manager.capabilities.dispatch(
            action=ActionRequest("slow", {}, "slow-call"),
            spec=self.main._contract()[0]["slow"],
            agent=self.main,
        )
        self.assertTrue(runtime._wait(marker.exists, timeout=5))
        self.stop_session()
        context = read_json(
            self.store.sessions / "session_000001/instances/main/context.json"
        )
        results = [
            json.loads(m["content"])
            for m in context
            if m["role"] == "user"
            and isinstance(m.get("content"), str)
            and "slow-call" in m["content"]
        ]
        self.assertEqual(len(results), 1)
        self.assertEqual(results[0]["data"]["status"], "interrupted")
        self.assertEqual(
            results[0]["data"]["info"], "Action interrupted by session shutdown."
        )

    def test_execution_process_death_delivers_an_interrupted_result(self):
        from jarvis.core.protocol import ActionRequest

        self.fixture.write_action(
            "suicide",
            runtime.EMPTY_SCHEMA,
            runtime.EMPTY_SCHEMA,
            "    import os, signal\n    os.kill(os.getpid(), signal.SIGKILL)\n",
        )
        self.manager.enable_action("main", "suicide")
        self.manager.capabilities.dispatch(
            action=ActionRequest("suicide", {}, "dead-call"),
            spec=self.main._contract()[0]["suicide"],
            agent=self.main,
        )
        self.assertTrue(
            runtime._wait(lambda: "dead-call" in str(self.main.history), timeout=5)
        )
        result = json.loads(
            next(
                m["content"]
                for m in self.main.history
                if m["role"] == "user" and "dead-call" in str(m)
            )
        )
        self.assertEqual(
            result["data"],
            {
                "status": "interrupted",
                "info": "Action interrupted because its execution process was terminated.",
            },
        )

    def test_resume_after_archive_commit_does_not_repeat_api_requests(self):
        self.main.enqueue(Event("tick.event", {"text": "Preserve committed archive"}))
        with patch.object(
            self.store,
            "clear_after_archive",
            side_effect=OSError("simulated cleanup interruption"),
        ):
            self.stop_session()
        self.assertEqual(self.store.state()["shutdown"]["stage"], "cleanup")
        archive = self.store.sessions / "session_000001/instances/main/context.json"
        original = archive.read_bytes()
        self.manager.shutdown()
        self.api = FakeMemoryAPI()
        self.manager = self.make_manager()
        self.manager.restoring_shutdown = True
        self.manager.frozen.set()
        self.main = self.manager.restore(preset="main")
        self.manager.services["session_restart_event"] = Event(
            "system_restarted", {"datetime": "2026-10-21T12:00:00+03:00"}, target="main"
        )
        self.stop_session()
        self.assertEqual(self.api.calls, [])
        self.assertEqual(archive.read_bytes(), original)
        final_input = self.manager.client.chat.completions.calls[-1]["messages"]
        self.assertTrue(
            any("system_restarted" in str(m.get("content")) for m in final_input)
        )
        self.assertNotIn("system_restarted", archive.read_text())
        self.assertEqual(read_json(self.store.instances / "main/context.json"), [])

    def test_successful_media_archive_returns_original_and_uses_api_profiles(self):
        image = {
            "type": "image_url",
            "image_url": {
                "url": "data:image/png;base64," + base64.b64encode(PNG).decode(),
                "name": "screenshot.png",
            },
        }
        self.main.history.append(
            {
                "role": "user",
                "content": [
                    {"type": "text", "text": '{"event_id":"picture","data":{}}'},
                    image,
                ],
            }
        )
        self.main._persist_context()
        self.stop_session()
        reader = self.manager.archive_reader
        filename, native = reader.get_file("main", "session_000001", "file_000001")
        self.assertEqual(filename, "screenshot.png")
        self.assertEqual(
            base64.b64decode(native["image_url"]["url"].split(",", 1)[1]), PNG
        )
        context = read_json(
            self.store.sessions / "session_000001/instances/main/context.json"
        )
        part = context[0]["content"][1]
        self.assertIn("Работа сохранена.", part["text_representation"])
        self.assertEqual(part["file_id"], "file_000001")
        self.assertEqual(
            [c["model"] for c in self.api.calls], ["api/image", "api/image", "api/text"]
        )
        self.assertEqual(self.api.calls[1]["messages"][1]["role"], "user")
        self.assertEqual(
            self.api.calls[1]["messages"][1]["content"][1]["type"], "image_url"
        )
        self.assertNotIn("base64", str(context))

    def test_other_final_action_is_not_executed_and_archive_remains_unchanged(self):
        self.manager.client = runtime._Client(
            [
                '{"actions":[{"action_id":"agents.spawn","call_id":"forbidden","data":{"preset_id":"worker"}}]}'
            ]
        )
        with patch.object(
            self.manager, "spawn", side_effect=AssertionError("Must not run")
        ):
            self.stop_session()
        self.assertTrue(self.manager.completed.is_set())
        self.assertEqual(self.store.state()["shutdown"]["stage"], "finished")
        self.assertNotIn(
            "forbidden",
            (
                self.store.sessions / "session_000001/instances/main/context.json"
            ).read_text(),
        )

    def test_runtime_shutdown_records_interruption_without_completing_logical_session(
        self,
    ):
        action = {
            "action_id": "say",
            "call_id": "runtime-stop",
            "data": {"text": "pending"},
        }
        self.main.history.append(
            {"role": "assistant", "content": json.dumps({"actions": [action]})}
        )
        self.main._persist_context()
        self.manager.results.begin(
            agent_id="main",
            call_id="runtime-stop",
            action_id="say",
            result_schema={"properties": {}},
        )
        self.manager.shutdown()
        self.assertIsNone(self.store.state()["shutdown"])
        self.assertFalse(self.store.sessions.exists())
        context = read_json(self.store.instances / "main/context.json")
        self.assertIn(
            "Action interrupted because its execution process was terminated.",
            str(context),
        )

    def test_incoming_input_is_durable_before_inference_and_thinking_survives(self):
        event = Event("tick.event", {"text": "Durable immediately"})
        self.main.enqueue(event)
        self.assertIn(
            "Durable immediately",
            (self.store.instances / "main/context.json").read_text(),
        )
        self.assertEqual(self.manager.client.chat.completions.calls, [])

        class Stream:
            def __iter__(self):
                yield SimpleNamespace(
                    choices=[
                        SimpleNamespace(
                            delta=SimpleNamespace(
                                content="",
                                refusal=None,
                                reasoning_content="Рассуждение",
                            )
                        )
                    ]
                )
                yield SimpleNamespace(
                    choices=[
                        SimpleNamespace(
                            delta=SimpleNamespace(
                                content='{"actions":[{"action_id":"no_action","call_id":"thinking-done","data":{}}]}',
                                refusal=None,
                            )
                        )
                    ]
                )

            def close(self):
                pass

        self.manager.client.chat.completions.create = lambda **kwargs: Stream()
        self.main._turn([event], self.main._generation)
        saved = read_json(self.store.instances / "main/context.json")
        self.assertEqual(saved[-1]["thinking"], "Рассуждение")
        self.assertEqual(len([m for m in saved if m["role"] == "user"]), 1)

    def test_main_shutdown_action_runs_complete_protocol_through_normal_dispatch(self):
        self.assertEqual(self.store.claim_start_event(), "system_started")
        self.main.start()
        self.manager.client = runtime._Client(
            [
                '{"actions":[{"action_id":"shutdown","call_id":"shutdown-via-model","data":{}}]}',
                '{"actions":[{"action_id":"reply","call_id":"farewell","data":{"text":"Готово."}}]}',
            ]
        )
        self.main.enqueue(Event("tick.event", {"text": "Заверши сессию"}))
        self.assertTrue(self.manager.completed.wait(10))
        self.assertEqual(len(self.manager.client.chat.completions.calls), 2)
        archived = read_json(
            self.store.sessions / "session_000001/instances/main/context.json"
        )
        self.assertIn("shutdown-via-model", str(archived))
        self.assertNotIn("farewell", str(archived))
        self.assertEqual(read_json(self.store.instances / "main/context.json"), [])
        self.assertFalse(self.store.state()["started"])
        self.assertEqual(self.store.claim_start_event(), "system_started")


class DocumentTests(unittest.TestCase):
    setUp = SessionMemoryTests.setUp
    record = SessionMemoryTests.record

    def image(self):
        from PIL import Image

        stream = io.BytesIO()
        Image.new("RGB", (8, 8), "red").save(stream, format="PNG")
        return stream.getvalue()

    def convert(self, path):
        self.store.save(self.record())
        item = self.store.register_file(
            "main", path.read_bytes(), path.name, "file", "application/octet-stream"
        )
        self.store.save(self.record(messages=[{"role": "user", "content": [item]}]))
        body = Documents(self.store, "main", "session_000001").convert(item)
        self.store.update_file("main", item)
        images = [f for f in self.store.files("main") if f["type"] == "image_url"]
        self.assertTrue(images)
        self.assertIn("memory://main/session_000001/", body)
        return body, images

    def test_docx_keeps_order_and_original_embedded_image(self):
        from docx import Document

        image = self.image()
        document = Document()
        document.add_paragraph("Before image")
        document.add_picture(io.BytesIO(image))
        document.add_paragraph("After image")
        path = self.root / "report.docx"
        document.save(path)
        body, images = self.convert(path)
        self.assertLess(body.index("Before image"), body.index(images[0]["file_id"]))
        self.assertLess(body.index(images[0]["file_id"]), body.index("After image"))
        self.assertEqual(
            (self.store.instances / "main/files" / images[0]["filename"]).read_bytes(),
            image,
        )

    def test_pptx_keeps_shape_order_and_original_embedded_image(self):
        from pptx import Presentation
        from pptx.util import Inches

        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.shapes.add_textbox(0, 0, Inches(4), Inches(1)).text = "Before picture"
        slide.shapes.add_picture(io.BytesIO(self.image()), 0, Inches(1))
        slide.shapes.add_textbox(
            0, Inches(2), Inches(4), Inches(1)
        ).text = "After picture"
        path = self.root / "slides.pptx"
        presentation.save(path)
        body, images = self.convert(path)
        self.assertLess(body.index("Before picture"), body.index(images[0]["file_id"]))
        self.assertLess(body.index(images[0]["file_id"]), body.index("After picture"))

    def test_pdf_interleaves_text_and_images_in_page_order(self):
        import pymupdf

        document = pymupdf.open()
        page = document.new_page()
        page.insert_text((40, 40), "Before PDF image")
        page.insert_image(pymupdf.Rect(40, 80, 100, 140), stream=self.image())
        page.insert_text((40, 200), "After PDF image")
        path = self.root / "document.pdf"
        document.save(path)
        document.close()
        body, images = self.convert(path)
        self.assertLess(
            body.index("Before PDF image"), body.index(images[0]["file_id"])
        )
        self.assertLess(body.index(images[0]["file_id"]), body.index("After PDF image"))

    def test_xlsx_saves_images_and_preserves_cell_values(self):
        import openpyxl
        from openpyxl.drawing.image import Image

        book = openpyxl.Workbook()
        sheet = book.active
        sheet["A1"] = "Before spreadsheet image"
        sheet.add_image(Image(io.BytesIO(self.image())), "A2")
        sheet["A3"] = "After spreadsheet image"
        path = self.root / "table.xlsx"
        book.save(path)
        body, images = self.convert(path)
        self.assertLess(
            body.index("Before spreadsheet image"), body.index(images[0]["file_id"])
        )
        self.assertLess(
            body.index(images[0]["file_id"]), body.index("After spreadsheet image")
        )

    def test_html_embedded_data_uri_gets_its_own_original_file(self):
        path = self.root / "page.html"
        path.write_text(
            '<h1>Заголовок</h1><p>Before</p><img src="data:image/png;base64,'
            + base64.b64encode(self.image()).decode()
            + '"><p>After</p>'
        )
        body, images = self.convert(path)
        self.assertIn("# Заголовок", body)
        self.assertLess(body.index("Before"), body.index(images[0]["file_id"]))
        self.assertLess(body.index(images[0]["file_id"]), body.index("After"))
